#!/usr/bin/env python3
"""
HJTC Verkaufs-Sync
==================
Holt bezahlte eBay-Bestellungen und trägt in der Buchhaltung (OneDrive, Buchhaltung_HJTC_v2.xlsx,
Blatt „Buchhaltung“) die echten Werte ein – sobald du die Zeile auf „Verkauft“ gesetzt hast:

  O  Verkaufsdatum      (nur wenn leer)
  F  Käuferversand      (was der Käufer laut eBay bezahlt hat, Versandrabatt berücksichtigt, bei Mehrfachbestellung anteilig)
  K  eBay-Gebühr        (tatsächliche Gebühr der Bestellung, anteilig nach Artikelwert; ersetzt die Formel)
  T  eBay-Artikelnr.

Zuordnung Bestellung -> Zeile: zuerst über die eBay-Artikelnr. in Spalte T, sonst über Titelähnlichkeit + Preis.
Steht die Zeile noch nicht auf „Verkauft“, wartet der Sync (du setzt den Status selbst) und versucht es beim nächsten Lauf.
Bereits übertragene Bestellpositionen merkt er sich (sync_state.json) und fasst sie nie wieder an.

Schreibzugriff auf Excel über Microsoft Graph (Excel im Web bleibt dabei offen nutzbar).

Aufruf:
  python verkaufs_sync.py --einrichten          einmalig: Microsoft-Anmeldung (Gerätecode)
  python verkaufs_sync.py --test                Probelauf: zeigt, was eingetragen würde, schreibt nichts
  python verkaufs_sync.py --test --lokal DATEI  Probelauf gegen eine lokale Kopie (ohne Microsoft)
  python verkaufs_sync.py                       normaler Lauf (systemd-Timer, stündlich)
"""
import os, sys, re, json, time, base64, logging, difflib
from pathlib import Path
from datetime import datetime, timedelta, timezone, date
import requests

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(Path.home() / 'ebay-status'))
try:
    import ereignisse                    # gemeinsames Ereignis-Protokoll (Morgenbericht/Cockpit)
except ImportError:
    ereignisse = None
STATE = HERE / 'sync_state.json'
TOKEN_CACHE = HERE / 'ms_token_cache.json'

DATEI = 'Buchhaltung_HJTC_v2.xlsx'          # im OneDrive-Hauptordner
BLATT = 'Buchhaltung'
LETZTE_ZEILE = 999
START = os.environ.get('SYNC_START', '2026-10-21')                       # nur Bestellungen ab diesem Tag (Altbestand ist bereits gepflegt)
TAGE_ZURUECK = 45                            # so weit zurück Bestellungen holen
ERINNERN_NACH_TAGEN = 3                      # Hinweis, wenn Status nach X Tagen noch nicht „Verkauft“
MS_SCOPES = ['Files.ReadWrite', 'User.Read']   # offline_access ergänzt MSAL selbst
# Spalten (1-basiert): A Kategorie, B Karte, C Status, E VK, F Versand, K Gebühr, O Verkaufsdatum, T eBay-Nr.
SP = {'B': 2, 'C': 3, 'E': 5, 'F': 6, 'K': 11, 'O': 15, 'T': 20}

logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s')
log = logging.getLogger('sync')


# ---------------------------------------------------------------- Zustand
def lade_state():
    try:
        return json.loads(STATE.read_text())
    except (FileNotFoundError, json.JSONDecodeError):
        return {'erledigt': {}, 'erinnert': {}}


def speichere_state(st):
    tmp = STATE.with_suffix('.tmp')
    tmp.write_text(json.dumps(st, indent=1, ensure_ascii=False))
    tmp.replace(STATE)


# ---------------------------------------------------------------- eBay
def ebay_token():
    creds = base64.b64encode(f"{os.environ['EBAY_CLIENT_ID']}:{os.environ['EBAY_CLIENT_SECRET']}".encode()).decode()
    r = requests.post('https://api.ebay.com/identity/v1/oauth2/token', timeout=20,
                      headers={'Content-Type': 'application/x-www-form-urlencoded', 'Authorization': f'Basic {creds}'},
                      data={'grant_type': 'refresh_token', 'refresh_token': os.environ['EBAY_REFRESH_TOKEN'],
                            'scope': 'https://api.ebay.com/oauth/api_scope https://api.ebay.com/oauth/api_scope/sell.fulfillment'})
    r.raise_for_status()
    return r.json()['access_token']


def betrag(x):
    return float((x or {}).get('value') or 0)


def bestellpositionen():
    """Bezahlte, nicht stornierte Bestellungen -> eine Position je Artikel mit anteiligem Versand und Gebühr."""
    tok = ebay_token()
    von = max(datetime.fromisoformat(START), datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(days=TAGE_ZURUECK))
    out, offset = [], 0
    if von > datetime.now(timezone.utc).replace(tzinfo=None):
        return out                     # Startdatum liegt noch in der Zukunft
    while True:
        r = requests.get('https://api.ebay.com/sell/fulfillment/v1/order', timeout=30,
                         headers={'Authorization': f'Bearer {tok}', 'X-EBAY-C-MARKETPLACE-ID': 'EBAY_DE'},
                         params={'filter': f'creationdate:[{von:%Y-%m-%dT00:00:00.000Z}..]', 'limit': 200, 'offset': offset})
        r.raise_for_status()
        d = r.json()
        for o in d.get('orders', []):
            if o.get('orderPaymentStatus') != 'PAID' or (o.get('cancelStatus') or {}).get('cancelState') not in (None, 'NONE_REQUESTED'):
                continue
            p = o.get('pricingSummary') or {}
            versand_gesamt = max(betrag(p.get('deliveryCost')) + betrag(p.get('deliveryDiscount')), 0.0)
            gebuehr = betrag(o.get('totalMarketplaceFee'))
            lines = o.get('lineItems', [])
            wert = [betrag(li.get('lineItemCost')) for li in lines]
            summe = sum(wert) or 1.0
            for li, w in zip(lines, wert):
                # Versand (nach Rabatt) und Gebühr der Bestellung anteilig nach Artikelwert verteilen
                anteil = w / summe if sum(wert) else 1 / len(lines)
                versand = versand_gesamt * anteil
                out.append({'schluessel': f"{o['orderId']}/{li.get('lineItemId')}", 'order': o['orderId'],
                            'item_id': li.get('legacyItemId'), 'titel': li.get('title', ''), 'preis': w,
                            'menge': li.get('quantity', 1), 'versand': round(versand, 2), 'gebuehr': round(gebuehr * anteil, 2),
                            'datum': o['creationDate'][:10], 'erstattet': bool((o.get('paymentSummary') or {}).get('refunds'))})
        offset += 200
        if offset >= d.get('total', 0):
            break
    return out


# ---------------------------------------------------------------- Excel: Microsoft Graph oder lokale Datei
class GraphExcel:
    BASIS = 'https://graph.microsoft.com/v1.0'

    def __init__(self):
        import msal
        self.cache = msal.SerializableTokenCache()
        if TOKEN_CACHE.exists():
            self.cache.deserialize(TOKEN_CACHE.read_text())
        self.app = msal.PublicClientApplication(os.environ['MS_CLIENT_ID'], authority='https://login.microsoftonline.com/consumers',
                                                token_cache=self.cache)
        self.session = None

    def _cache_sichern(self):
        if self.cache.has_state_changed:
            TOKEN_CACHE.write_text(self.cache.serialize())
            TOKEN_CACHE.chmod(0o600)

    def einrichten(self):
        flow = self.app.initiate_device_flow(scopes=MS_SCOPES)
        if 'user_code' not in flow:
            raise RuntimeError(f'Gerätecode fehlgeschlagen: {flow}')
        print('\n' + flow['message'] + '\n', flush=True)
        res = self.app.acquire_token_by_device_flow(flow)
        if 'access_token' not in res:
            raise RuntimeError(f"Anmeldung fehlgeschlagen: {res.get('error_description')}")
        self._cache_sichern()
        print('Anmeldung gespeichert.')

    def token(self):
        konten = self.app.get_accounts()
        res = self.app.acquire_token_silent(MS_SCOPES, account=konten[0]) if konten else None
        if not res or 'access_token' not in res:
            raise RuntimeError('Keine gültige Microsoft-Anmeldung. Bitte einmal „--einrichten“ ausführen.')
        self._cache_sichern()
        return res['access_token']

    def _h(self):
        h = {'Authorization': f'Bearer {self.token()}', 'Content-Type': 'application/json'}
        if self.session:
            h['workbook-session-id'] = self.session
        return h

    def _wb(self):
        return f"{self.BASIS}/me/drive/root:/{requests.utils.quote(DATEI)}:/workbook"

    def lesen(self):
        if not self.session:
            r = requests.post(f'{self._wb()}/createSession', headers=self._h(), json={'persistChanges': True}, timeout=60)
            r.raise_for_status()
            self.session = r.json()['id']
        r = requests.get(f"{self._wb()}/worksheets('{BLATT}')/range(address='A1:X{LETZTE_ZEILE}')?$select=values,formulas",
                         headers=self._h(), timeout=60)
        r.raise_for_status()
        j = r.json()
        return j['values'], j['formulas']

    def schreiben(self, zeile, spalte, wert):
        adr = f'{"ABCDEFGHIJKLMNOPQRSTUVWX"[spalte - 1]}{zeile}'
        r = requests.patch(f"{self._wb()}/worksheets('{BLATT}')/range(address='{adr}')", headers=self._h(),
                           json={'values': [[wert]]}, timeout=60)
        r.raise_for_status()

    def schliessen(self):
        if self.session:
            requests.post(f'{self._wb()}/closeSession', headers=self._h(), timeout=30)
            self.session = None


class LokalExcel:
    """Nur für Probeläufe: liest eine lokale Kopie, schreibt nichts."""
    def __init__(self, pfad):
        self.pfad = pfad

    def lesen(self):
        import openpyxl
        wf = openpyxl.load_workbook(self.pfad)[BLATT]
        wv = openpyxl.load_workbook(self.pfad, data_only=True)[BLATT]
        vals, forms = [], []
        for r in range(1, LETZTE_ZEILE + 1):
            vals.append([_excel_wert(wv.cell(r, c).value) for c in range(1, 25)])
            forms.append([wf.cell(r, c).value if wf.cell(r, c).value is not None else '' for c in range(1, 25)])
        return vals, forms

    def schreiben(self, zeile, spalte, wert):
        raise RuntimeError('lokaler Modus schreibt nicht')

    def schliessen(self):
        pass


def _excel_wert(v):
    if isinstance(v, datetime):
        return (v - datetime(1899, 12, 30)).days
    return '' if v is None else v


def excel_datum(iso):
    return (date.fromisoformat(iso) - date(1899, 12, 30)).days


# ---------------------------------------------------------------- Zuordnung
def norm(s):
    return ' '.join(re.findall(r'[a-z0-9]+', str(s).lower().replace('kakawow', '').replace('trading card', '')))


def zuordnen(pos, zeilen, belegt):
    """Gibt (zeilennummer, grund) zurück oder (None, grund)."""
    for nr, z in zeilen.items():
        if str(z['T']).strip() == str(pos['item_id']) and nr not in belegt:
            return nr, 'eBay-Artikelnr.'
    kandidaten = []
    for nr, z in zeilen.items():
        if nr in belegt or not z['B']:
            continue
        if str(z['T']).strip() and str(z['T']).strip() != str(pos['item_id']):
            continue                       # gehört nachweislich zu einem anderen Angebot
        s = difflib.SequenceMatcher(None, norm(pos['titel']), norm(z['B'])).ratio()
        try:
            preis_ok = abs(float(z['E']) - pos['preis']) <= max(1.0, 0.15 * pos['preis'])
        except (TypeError, ValueError):
            preis_ok = False
        if s >= 0.55 and preis_ok:
            kandidaten.append((s, nr))
    kandidaten.sort(reverse=True)
    if not kandidaten:
        return None, 'keine passende Zeile'
    if len(kandidaten) > 1 and kandidaten[0][0] - kandidaten[1][0] < 0.08:
        return None, f'mehrdeutig (Zeilen {kandidaten[0][1]} und {kandidaten[1][1]})'
    return kandidaten[0][1], f'Titel {kandidaten[0][0]:.0%} + Preis'


# ---------------------------------------------------------------- Lauf
def telegram(text, test):
    if test:
        print('\n--- TELEGRAM ---\n' + text)
        return
    requests.post(f"https://api.telegram.org/bot{os.environ['TELEGRAM_BOT_TOKEN']}/sendMessage", timeout=30,
                  data={'chat_id': os.environ['TELEGRAM_CHAT_ID'], 'text': text[:4000], 'disable_web_page_preview': 'true'})


def lauf(excel, test):
    st = lade_state()
    positionen = [p for p in bestellpositionen() if p['schluessel'] not in st['erledigt']]
    log.info('%s neue Bestellpositionen seit %s', len(positionen), START)
    if not positionen:
        return
    werte, formeln = excel.lesen()
    zeilen = {i + 1: {k: werte[i][c - 1] for k, c in SP.items()} | {'K_formel': formeln[i][SP['K'] - 1]}
              for i in range(1, len(werte)) if werte[i][SP['B'] - 1]}
    belegt, eingetragen, wartend, offen = set(), [], [], []
    for p in sorted(positionen, key=lambda x: x['datum']):
        nr, grund = zuordnen(p, zeilen, belegt)
        if nr is None:
            if f"offen:{p['schluessel']}" not in st['erinnert']:          # nur einmal melden
                offen.append(f"• {p['titel'][:60]} ({p['datum']}, {p['preis']:.2f} €): {grund}")
                if not test:
                    st['erinnert'][f"offen:{p['schluessel']}"] = date.today().isoformat()
            continue
        z = zeilen[nr]
        if str(z['C']).strip() not in ('Verkauft', 'Verkauft (Lot)'):
            alter = (date.today() - date.fromisoformat(p['datum'])).days
            if alter >= ERINNERN_NACH_TAGEN and p['schluessel'] not in st['erinnert']:
                wartend.append(f"• Zeile {nr}: {str(z['B'])[:55]} (verkauft am {p['datum']})")
                if not test:
                    st['erinnert'][p['schluessel']] = date.today().isoformat()
            continue
        aenderungen = []
        if not z['O']:
            aenderungen.append(('O', excel_datum(p['datum']), p['datum']))
        aenderungen += [('F', p['versand'], f"{p['versand']:.2f} €"), ('K', p['gebuehr'], f"{p['gebuehr']:.2f} €")]
        if str(z['T']).strip() != str(p['item_id']):
            aenderungen.append(('T', str(p['item_id']), str(p['item_id'])))
        if not test:
            for sp, wert, _ in aenderungen:
                excel.schreiben(nr, SP[sp], wert)
            st['erledigt'][p['schluessel']] = {'zeile': nr, 'datum': date.today().isoformat()}
            speichere_state(st)
        belegt.add(nr)
        eingetragen.append(f"• Zeile {nr}: {str(z['B'])[:50]} · " + ', '.join(f'{sp} {txt}' for sp, _, txt in aenderungen)
                           + f" ({grund}){' · ⚠ Erstattung' if p['erstattet'] else ''}")
    excel.schliessen()
    if not test:
        speichere_state(st)
    if ereignisse and not test:
        for z in eingetragen:
            ereignisse.melden('Verkaufs-Sync', 'info', 'Eingetragen: ' + z.lstrip('• '))
        for z in wartend:
            ereignisse.melden('Verkaufs-Sync', 'aktion', 'Status auf „Verkauft“ setzen: ' + z.lstrip('• '))
        for z in offen:
            ereignisse.melden('Verkaufs-Sync', 'aktion', 'eBay-Nr. in Spalte T eintragen: ' + z.lstrip('• '))
        log.info('eingetragen %s, wartend %s, offen %s (gemeldet im Morgenbericht)', len(eingetragen), len(wartend), len(offen))
        return
    teile = []
    if eingetragen:
        teile.append('📒 Verkaufs-Sync: in der Buchhaltung ergänzt\n' + '\n'.join(eingetragen))
    if wartend:
        teile.append('⏳ Verkauft, aber Status noch nicht „Verkauft“ – bitte setzen, dann trage ich den Rest ein:\n' + '\n'.join(wartend))
    if offen:
        teile.append('❓ Keine eindeutige Zeile gefunden (bitte eBay-Artikelnr. in Spalte T eintragen):\n' + '\n'.join(offen))
    if teile:
        telegram('\n\n'.join(teile), test)
    log.info('eingetragen %s, wartend %s, offen %s', len(eingetragen), len(wartend), len(offen))


def main():
    test = '--test' in sys.argv
    if '--einrichten' in sys.argv:
        GraphExcel().einrichten()
        return
    if '--lokal' in sys.argv:
        excel = LokalExcel(sys.argv[sys.argv.index('--lokal') + 1])
        test = True
    else:
        excel = GraphExcel()
    lauf(excel, test)


if __name__ == '__main__':
    main()
